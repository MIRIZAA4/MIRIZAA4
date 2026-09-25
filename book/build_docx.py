# -*- coding: utf-8 -*-
"""בונה קובץ Word מאוחד מכל הפרקים בתיקיית chapters."""
import glob, os, re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = "בת הכפר"

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)

def add_par(doc, text, size=13, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, indent=True):
    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.3
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        pf.first_line_indent = Cm(0.6)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'David'
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:cs'), 'David'); rFonts.set(qn('w:ascii'), 'David'); rFonts.set(qn('w:hAnsi'), 'David')
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)
    return p

def main():
    doc = Document()
    for s in doc.sections:
        s.page_height = Cm(21); s.page_width = Cm(14.8)   # A5
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.8)
    add_par(doc, TITLE, size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30, indent=False)
    add_par(doc, "רומן", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    doc.add_page_break()
    files = sorted(glob.glob(os.path.join(HERE, 'chapters', '*.md')))
    total = 0
    for i, f in enumerate(files):
        text = open(f, encoding='utf-8').read().strip()
        total += len(text.split())
        if i > 0:
            doc.add_page_break()
        for line in text.split('\n'):
            line = line.rstrip()
            if not line:
                continue
            if line.startswith('# '):
                add_par(doc, line[2:], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18, indent=False)
            elif line.strip() == '* * *':
                add_par(doc, '* * *', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, indent=False)
            else:
                add_par(doc, line)
    out = os.path.join(HERE, TITLE + '.docx')
    doc.save(out)
    print(f"{len(files)} פרקים, {total} מילים -> {out}")

if __name__ == '__main__':
    main()
